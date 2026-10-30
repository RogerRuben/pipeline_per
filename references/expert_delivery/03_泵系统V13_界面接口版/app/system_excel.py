# -*- coding: utf-8 -*-
"""Generic system-level V11 reader for the V10-compatible Excel contract."""

from __future__ import annotations

import hashlib
import json
import math
import re
import shutil
import sys
from dataclasses import asdict, dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple, Union

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

from project_excel import AttributeSpec, ProjectDataError, ProjectDataset, load_project_workbook  # noqa: E402
from formula_constraints import FormulaConstraintSpec, SafeFormulaExpression  # noqa: E402


CONTRACT_VERSION = "system_excel_v10"
SYSTEM_ATTRIBUTE_SHEET = "系统属性配置"
MAPPING_SHEET = "系统映射配置"
TASK_SHEET = "系统任务要求"
FORMULA_SHEET = "派生属性公式"

SYSTEM_ATTRIBUTE_HEADERS = [
    "系统属性",
    "系统属性ID",
    "系统单位",
    "汇总方式",
    "说明",
]

MAPPING_HEADERS = [
    "映射编号",
    "成品属性",
    "属性ID",
    "评价通道",
    "系统属性",
    "系统属性ID",
    "映射类型",
    "映射系数",
    "映射常数",
    "说明",
]

BASE_TASK_HEADERS = [
    "任务编号",
    "任务名称",
    "系统型号/工况",
    "系统属性",
    "系统属性ID",
    "要求类型",
    "目标值",
    "要求下限",
    "要求上限",
    "下偏80分距离",
    "上偏80分距离",
    "允许超额加分",
    "单项分数上限",
    "组内权重",
    "硬性策略",
    "说明",
]

TASK_HEADERS = [
    "任务编号",
    "任务名称",
    "系统型号/工况",
    "系统属性",
    "系统属性ID",
    "要求类型",
    "目标值",
    "要求下限",
    "要求上限",
    "下偏80分距离",
    "上偏80分距离",
    "达标允许下偏差",
    "达标允许上偏差",
    "允许超额加分",
    "单项分数上限",
    "组内权重",
    "硬性策略",
    "说明",
]

FORMULA_HEADERS = [
    "公式编号",
    "目标属性",
    "目标属性ID",
    "公式表达式",
    "关系类型",
    "处理模式",
    "相对容差",
    "绝对容差",
    "约束级别",
    "生成时自动回算",
    "适用工况",
    "说明",
]

LEGACY_FORMULA_HEADERS = [
    "公式编号", "目标属性", "目标属性ID", "公式类型", "输入属性ID", "换算常数",
    "处理模式", "相对容差", "绝对容差", "约束级别", "生成时自动回算", "说明",
]

CHANNEL_ALIASES = {
    "直接系统贡献": "direct_system",
    "direct_system": "direct_system",
    "成品性能贡献": "product_performance",
    "product_performance": "product_performance",
    "仅可行性": "feasibility_only",
    "feasibility_only": "feasibility_only",
}

MAPPING_TYPE_ALIASES = {
    "一比一": "linear",
    "线性传导": "linear",
    "identity": "linear",
    "linear": "linear",
    "加总贡献": "additive",
    "additive": "additive",
    "成品性能": "product_performance",
    "product_performance": "product_performance",
    "无": "none",
    "none": "none",
}

AGGREGATION_ALIASES = {
    "单值": "single",
    "单一来源": "single",
    "single": "single",
    "求和": "sum",
    "加总": "sum",
    "sum": "sum",
}

REQUIREMENT_TYPE_ALIASES = {
    "目标值": "target",
    "接近目标值最好": "target",
    "target": "target",
    "下限": "at_least",
    "不低于": "at_least",
    "at_least": "at_least",
    "上限/预算": "at_most",
    "上限": "at_most",
    "不超过": "at_most",
    "at_most": "at_most",
    "区间内满足": "within_range",
    "区间要求": "within_range",
    "within_range": "within_range",
    "相对越大越好": "relative_higher",
    "relative_higher": "relative_higher",
    "相对越小越好": "relative_lower",
    "relative_lower": "relative_lower",
}

HARD_POLICY_ALIASES = {
    "否": "none",
    "无": "none",
    "none": "none",
    "一票否决": "veto",
    "veto": "veto",
    "限制总分": "cap_total",
    "cap_total": "cap_total",
}

FORMULA_TYPE_ALIASES = {
    "乘积除常数": "product_divide_constant",
    "压力×流量÷常数": "product_divide_constant",
    "压力*流量/常数": "product_divide_constant",
    "product_divide_constant": "product_divide_constant",
}

FORMULA_MODE_ALIASES = {
    "严格等式": "exact",
    "exact": "exact",
    "效率折算": "efficiency_calibrated",
    "样本折算": "efficiency_calibrated",
    "efficiency_calibrated": "efficiency_calibrated",
    "比例折算": "scale_calibrated",
    "样本比例折算": "scale_calibrated",
    "scale_calibrated": "scale_calibrated",
    "仿射折算": "affine_calibrated",
    "线性校准": "affine_calibrated",
    "affine_calibrated": "affine_calibrated",
    "仅诊断": "diagnostic",
    "diagnostic": "diagnostic",
}

FORMULA_RELATION_ALIASES = {
    "等于": "equal",
    "等式": "equal",
    "=": "equal",
    "equal": "equal",
    "不超过": "less_equal",
    "小于等于": "less_equal",
    "<=": "less_equal",
    "less_equal": "less_equal",
    "不低于": "greater_equal",
    "大于等于": "greater_equal",
    ">=": "greater_equal",
    "greater_equal": "greater_equal",
}

FORMULA_ENFORCEMENT_ALIASES = {
    "硬约束": "hard",
    "hard": "hard",
    "软约束": "soft",
    "soft": "soft",
    "仅提示": "diagnostic",
    "diagnostic": "diagnostic",
}


@dataclass(frozen=True)
class SystemAttributeSpec:
    key: str
    label: str
    unit: str
    aggregation: str = "single"
    description: str = ""


@dataclass(frozen=True)
class SystemMappingSpec:
    mapping_id: str
    product_key: str
    product_label: str
    product_unit: str
    channel: str
    system_key: str = ""
    system_attribute: str = ""
    system_unit: str = ""
    mapping_type: str = "none"
    coefficient: float = 1.0
    constant: float = 0.0
    description: str = ""


@dataclass(frozen=True)
class SystemRequirementSpec:
    system_key: str
    system_attribute: str
    system_unit: str
    requirement_type: str
    target_value: Optional[float] = None
    minimum: Optional[float] = None
    maximum: Optional[float] = None
    lower_80_distance: Optional[float] = None
    upper_80_distance: Optional[float] = None
    overachievement_bonus: bool = False
    score_cap: float = 100.0
    weight: float = 1.0
    hard_policy: str = "none"
    description: str = ""
    lower_compliance_tolerance: Optional[float] = None
    upper_compliance_tolerance: Optional[float] = None


@dataclass(frozen=True)
class SystemTask:
    task_id: str
    task_name: str
    operating_condition: str = ""
    requirements: Tuple[SystemRequirementSpec, ...] = ()
    description: str = ""

    def requirement_by_key(self) -> Dict[str, SystemRequirementSpec]:
        return {item.system_key: item for item in self.requirements}


@dataclass
class SystemProjectDataset:
    product: ProjectDataset
    system_attributes: List[SystemAttributeSpec]
    mappings: List[SystemMappingSpec]
    tasks: List[SystemTask]
    system_fingerprint: str
    system_learning_fingerprint: str
    formulas: List[FormulaConstraintSpec] = field(default_factory=list)
    contract_version: str = CONTRACT_VERSION
    source_contract_version: str = CONTRACT_VERSION
    warnings: List[str] = field(default_factory=list)

    def system_attribute_by_key(self) -> Dict[str, SystemAttributeSpec]:
        return {item.key: item for item in self.system_attributes}

    def mapping_by_key(self) -> Dict[str, List[SystemMappingSpec]]:
        output: Dict[str, List[SystemMappingSpec]] = {}
        for item in self.mappings:
            output.setdefault(item.product_key, []).append(item)
        return output

    def task_by_id(self) -> Dict[str, SystemTask]:
        return {item.task_id: item for item in self.tasks}

    def default_task(self) -> SystemTask:
        if not self.tasks:
            raise ProjectDataError("系统任务要求中没有可用任务。")
        return self.tasks[0]

    def direct_mappings(self) -> List[SystemMappingSpec]:
        return [item for item in self.mappings if item.channel == "direct_system"]

    def product_mappings(self) -> List[SystemMappingSpec]:
        return [item for item in self.mappings if item.channel == "product_performance"]

    def summary(self) -> Dict[str, Any]:
        return {
            "contract_version": self.contract_version,
            "source_contract_version": self.source_contract_version,
            "product": self.product.summary(),
            "system_fingerprint": self.system_fingerprint,
            "system_learning_fingerprint": self.system_learning_fingerprint,
            "system_attributes": [asdict(item) for item in self.system_attributes],
            "direct_mappings": [asdict(item) for item in self.direct_mappings()],
            "product_performance_attributes": [item.product_label for item in self.product_mappings()],
            "tasks": [asdict(item) for item in self.tasks],
            "formulas": [asdict(item) for item in self.formulas],
            "warnings": list(self.warnings),
        }


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _normal(value: Any) -> str:
    return re.sub(r"\s+", "", _text(value)).lower()


def _optional_float(value: Any, field_name: str) -> Optional[float]:
    if value is None or _text(value) == "":
        return None
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ProjectDataError(f"{field_name}必须是数值。") from exc
    if not math.isfinite(result):
        raise ProjectDataError(f"{field_name}必须是有限数值。")
    return result


def _bool(value: Any, default: bool = False) -> bool:
    if value is None or _text(value) == "":
        return default
    return _normal(value) in {"是", "true", "yes", "1", "y"}


def _enum(value: Any, aliases: Dict[str, str], field_name: str, default: str) -> str:
    key = _normal(value)
    if not key:
        return default
    if key not in aliases:
        choices = "、".join(sorted({item for item in aliases if not item.isascii()}))
        raise ProjectDataError(f"{field_name}填写无效，可填写：{choices}。")
    return aliases[key]


def _stable_key(value: Any, label: str, field_name: str) -> str:
    key = _text(value)
    if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", key):
        raise ProjectDataError(f"{field_name}“{label}”的ID {key!r} 无效，只能使用英文字母、数字和下划线。")
    return key


def _generated_key(label: str, prefix: str, index: int) -> str:
    ascii_key = re.sub(r"[^A-Za-z0-9_]+", "_", label).strip("_").lower()
    if ascii_key and ascii_key[0].isalpha():
        return ascii_key
    digest = hashlib.sha256(label.encode("utf-8")).hexdigest()[:8]
    return f"{prefix}_{index:02d}_{digest}"


def _header_map(sheet: Any) -> Dict[str, int]:
    output: Dict[str, int] = {}
    for index, cell in enumerate(sheet[1], start=1):
        name = _normal(cell.value)
        if name:
            output[name] = index
    return output


def _has_headers(sheet: Any, headers: Sequence[str]) -> bool:
    existing = _header_map(sheet)
    return all(_normal(item) in existing for item in headers)


def _cell(sheet: Any, row: int, headers: Dict[str, int], name: str) -> Any:
    column = headers.get(_normal(name))
    return sheet.cell(row=row, column=column).value if column else None


def _style_table(sheet: Any, widths: Sequence[int]) -> None:
    fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="16324F")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _keyword_score(spec: AttributeSpec, keywords: Sequence[str]) -> int:
    text = _normal(f"{spec.label} {spec.key} {spec.description}")
    score = 0
    for keyword in keywords:
        keyword = keyword.lower()
        if _normal(spec.label) == keyword or _normal(spec.key) == keyword:
            score += 20
        elif keyword in text:
            score += 5
    return score


def detect_direct_attributes(attributes: Sequence[AttributeSpec]) -> Dict[str, str]:
    """Suggest common pump roles only; the V2 contract never depends on these names."""
    groups = {
        "pressure": ("泵压力", "出口压力", "额定压力", "压力", "pressure", "p_out"),
        "flow": ("泵流量", "额定流量", "流量", "flow", "q_out"),
        "weight": ("泵重量", "产品重量", "重量", "质量", "weight", "mass"),
    }
    selected: Dict[str, str] = {}
    used: set[str] = set()
    for role, keywords in groups.items():
        ranked = sorted(
            ((_keyword_score(spec, keywords), spec.key) for spec in attributes if spec.is_numeric),
            reverse=True,
        )
        if ranked and ranked[0][0] > 0 and ranked[0][1] not in used:
            selected[role] = ranked[0][1]
            used.add(ranked[0][1])
    return selected


def _default_contract(product: ProjectDataset) -> Tuple[List[SystemAttributeSpec], List[SystemMappingSpec], List[SystemTask]]:
    detected = detect_direct_attributes(product.attributes)
    roles_by_key = {key: role for role, key in detected.items()}
    defaults = {
        "pressure": ("system_pressure", "系统压力", "single", "target", False, 100.0),
        "flow": ("system_flow", "系统流量", "single", "at_least", True, 130.0),
        "weight": ("system_weight_contribution", "系统重量贡献", "sum", "at_most", True, 130.0),
    }
    system_attributes: List[SystemAttributeSpec] = []
    mappings: List[SystemMappingSpec] = []
    requirements: List[SystemRequirementSpec] = []
    direct_keys: set[str] = set()
    mapping_index = 1
    for spec in product.attributes:
        role = roles_by_key.get(spec.key)
        if role:
            system_key, system_label, aggregation, req_type, bonus, cap = defaults[role]
            system_attributes.append(
                SystemAttributeSpec(system_key, system_label, spec.unit, aggregation, "自动识别结果，请由用户确认。")
            )
            mappings.append(
                SystemMappingSpec(
                    f"MAP-{mapping_index:03d}", spec.key, spec.label, spec.unit,
                    "direct_system", system_key, system_label, spec.unit,
                    "additive" if aggregation == "sum" else "linear", 1.0, 0.0,
                    "自动识别并预填；可在Excel中修改。",
                )
            )
            requirements.append(
                SystemRequirementSpec(
                    system_key, system_label, spec.unit, req_type,
                    overachievement_bonus=bonus, score_cap=cap,
                    description="请填写当前系统任务的要求数值。",
                )
            )
            direct_keys.add(spec.key)
        elif spec.participates_utility:
            mappings.append(
                SystemMappingSpec(
                    f"MAP-{mapping_index:03d}", spec.key, spec.label, spec.unit,
                    "product_performance", mapping_type="product_performance",
                    description="按成品属性方向进入经验型性能贡献组。",
                )
            )
        else:
            mappings.append(
                SystemMappingSpec(
                    f"MAP-{mapping_index:03d}", spec.key, spec.label, spec.unit,
                    "feasibility_only", mapping_type="none", description="仅参与物理可行性。",
                )
            )
        mapping_index += 1
    task = SystemTask("SYS-TASK-001", "系统任务1", "请填写系统型号或工况", tuple(requirements))
    return system_attributes, mappings, [task]


def _write_formula_sheet(workbook: Any, formulas: Sequence[FormulaConstraintSpec] = ()) -> None:
    if FORMULA_SHEET in workbook.sheetnames:
        del workbook[FORMULA_SHEET]
    sheet = workbook.create_sheet(FORMULA_SHEET)
    sheet.append(FORMULA_HEADERS)
    mode_labels = {
        "exact": "严格等式",
        "efficiency_calibrated": "效率折算",
        "scale_calibrated": "比例折算",
        "affine_calibrated": "仿射折算",
        "diagnostic": "仅诊断",
    }
    relation_labels = {"equal": "等于", "less_equal": "不超过", "greater_equal": "不低于"}
    enforcement_labels = {"hard": "硬约束", "soft": "软约束", "diagnostic": "仅提示"}
    for item in formulas:
        sheet.append(
            [
                item.formula_id,
                item.target_label,
                item.target_key,
                item.normalized_expression(),
                relation_labels[item.relation],
                mode_labels[item.mode],
                item.relative_tolerance,
                item.absolute_tolerance,
                enforcement_labels[item.enforcement],
                "是" if item.auto_complete else "否",
                item.operating_condition,
                item.description,
            ]
        )
    _style_table(sheet, [16, 20, 24, 52, 14, 18, 14, 14, 14, 18, 24, 58])


def _write_v2_contract(workbook: Any, attributes: Sequence[SystemAttributeSpec], mappings: Sequence[SystemMappingSpec], tasks: Sequence[SystemTask]) -> None:
    for name in (SYSTEM_ATTRIBUTE_SHEET, MAPPING_SHEET, TASK_SHEET):
        if name in workbook.sheetnames:
            del workbook[name]

    attribute_sheet = workbook.create_sheet(SYSTEM_ATTRIBUTE_SHEET)
    attribute_sheet.append(SYSTEM_ATTRIBUTE_HEADERS)
    for item in attributes:
        attribute_sheet.append([
            item.label, item.key, item.unit, "求和" if item.aggregation == "sum" else "单值", item.description,
        ])
    _style_table(attribute_sheet, [22, 26, 14, 14, 48])

    mapping_sheet = workbook.create_sheet(MAPPING_SHEET)
    mapping_sheet.append(MAPPING_HEADERS)
    channel_labels = {
        "direct_system": "直接系统贡献",
        "product_performance": "成品性能贡献",
        "feasibility_only": "仅可行性",
    }
    mapping_labels = {"linear": "线性传导", "additive": "加总贡献", "product_performance": "成品性能", "none": "无"}
    for item in mappings:
        mapping_sheet.append([
            item.mapping_id, item.product_label, item.product_key, channel_labels[item.channel],
            item.system_attribute, item.system_key, mapping_labels[item.mapping_type],
            item.coefficient, item.constant, item.description,
        ])
    _style_table(mapping_sheet, [16, 20, 24, 18, 22, 26, 16, 13, 13, 48])

    task_sheet = workbook.create_sheet(TASK_SHEET)
    task_sheet.append(TASK_HEADERS)
    requirement_labels = {
        "target": "目标值",
        "at_least": "下限",
        "at_most": "上限/预算",
        "within_range": "区间内满足",
        "relative_higher": "相对越大越好",
        "relative_lower": "相对越小越好",
    }
    hard_labels = {"none": "否", "veto": "一票否决", "cap_total": "限制总分"}
    for task in tasks:
        if not task.requirements:
            task_sheet.append([task.task_id, task.task_name, task.operating_condition, *([""] * 14), task.description])
            continue
        for requirement in task.requirements:
            task_sheet.append([
                task.task_id, task.task_name, task.operating_condition,
                requirement.system_attribute, requirement.system_key,
                requirement_labels[requirement.requirement_type],
                requirement.target_value, requirement.minimum, requirement.maximum,
                requirement.lower_80_distance, requirement.upper_80_distance,
                requirement.lower_compliance_tolerance, requirement.upper_compliance_tolerance,
                "是" if requirement.overachievement_bonus else "否",
                requirement.score_cap, requirement.weight, hard_labels[requirement.hard_policy],
                requirement.description or task.description,
            ])
    _style_table(task_sheet, [18, 24, 28, 22, 26, 18, 13, 13, 13, 18, 18, 18, 18, 16, 16, 14, 14, 48])
    _write_formula_sheet(workbook)


def _read_formulas(workbook: Any, product: ProjectDataset) -> Tuple[List[FormulaConstraintSpec], List[str]]:
    if FORMULA_SHEET not in workbook.sheetnames:
        return [], ["当前Excel没有“派生属性公式”表，V11仍可运行，但不会执行公式回算。"]
    sheet = workbook[FORMULA_SHEET]
    is_v5 = _has_headers(sheet, FORMULA_HEADERS)
    is_legacy = _has_headers(sheet, LEGACY_FORMULA_HEADERS)
    if not is_v5 and not is_legacy:
        raise ProjectDataError(f"“{FORMULA_SHEET}”表头不完整，请使用V11模板或先运行升级工具。")
    headers = _header_map(sheet)
    by_key = product.attribute_by_key()
    by_label = {_normal(item.label): item for item in product.attributes}
    formulas: List[FormulaConstraintSpec] = []
    warnings: List[str] = []
    seen_ids: set[str] = set()
    seen_targets: set[str] = set()
    for row in range(2, sheet.max_row + 1):
        formula_id = _text(_cell(sheet, row, headers, "公式编号"))
        target_key_raw = _text(_cell(sheet, row, headers, "目标属性ID"))
        target_label_raw = _text(_cell(sheet, row, headers, "目标属性"))
        if not formula_id and not target_key_raw and not target_label_raw:
            continue
        if not formula_id:
            raise ProjectDataError(f"“{FORMULA_SHEET}”第{row}行缺少公式编号。")
        if formula_id in seen_ids:
            raise ProjectDataError(f"公式编号“{formula_id}”重复。")
        target = by_key.get(target_key_raw) or by_label.get(_normal(target_label_raw))
        if target is None:
            raise ProjectDataError(f"“{FORMULA_SHEET}”第{row}行目标属性不存在。")
        if target.key in seen_targets:
            raise ProjectDataError(f"目标属性“{target.label}”配置了多条公式，当前版本只允许一个直接公式。")
        if is_v5:
            expression = _text(_cell(sheet, row, headers, "公式表达式"))
            if not expression:
                raise ProjectDataError(f"公式“{formula_id}”缺少公式表达式。")
            try:
                parsed = SafeFormulaExpression(expression)
            except ValueError as exc:
                raise ProjectDataError(f"公式“{formula_id}”无效：{exc}") from exc
            source_keys = tuple(sorted(parsed.variable_names))
            constant = 1.0
            formula_type = "expression"
        else:
            source_text = _text(_cell(sheet, row, headers, "输入属性ID"))
            source_keys = tuple(item for item in re.split(r"[,，、;；\s]+", source_text) if item)
            if len(source_keys) < 2:
                raise ProjectDataError(f"公式“{formula_id}”至少需要两个输入属性ID。")
            constant_value = _optional_float(
                _cell(sheet, row, headers, "换算常数"), f"公式“{formula_id}”换算常数"
            )
            if constant_value is None or constant_value <= 0:
                raise ProjectDataError(f"公式“{formula_id}”的换算常数必须大于0。")
            constant = float(constant_value)
            expression = f"({' * '.join(source_keys)}) / {constant:g}"
            formula_type = _enum(
                _cell(sheet, row, headers, "公式类型"), FORMULA_TYPE_ALIASES,
                f"公式“{formula_id}”公式类型", "product_divide_constant",
            )
            warnings.append(f"公式“{formula_id}”已从V4乘积除常数格式转换为V11安全表达式。")
        unknown = [key for key in source_keys if key not in by_key]
        if unknown:
            raise ProjectDataError(f"公式“{formula_id}”包含未知属性ID：{'、'.join(unknown)}")
        if target.key in source_keys:
            raise ProjectDataError(f"公式“{formula_id}”不能直接引用目标属性自身。")
        relative_tolerance = _optional_float(
            _cell(sheet, row, headers, "相对容差"), f"公式“{formula_id}”相对容差"
        )
        absolute_tolerance = _optional_float(
            _cell(sheet, row, headers, "绝对容差"), f"公式“{formula_id}”绝对容差"
        )
        spec = FormulaConstraintSpec(
            formula_id=formula_id,
            target_key=target.key,
            target_label=target.label,
            target_unit=target.unit,
            formula_type=formula_type,
            source_keys=source_keys,
            source_labels=tuple(by_key[key].label for key in source_keys),
            constant=float(constant),
            mode=_enum(
                _cell(sheet, row, headers, "处理模式"), FORMULA_MODE_ALIASES,
                f"公式“{formula_id}”处理模式", "exact",
            ),
            relative_tolerance=0.05 if relative_tolerance is None else float(relative_tolerance),
            absolute_tolerance=0.0 if absolute_tolerance is None else float(absolute_tolerance),
            enforcement=_enum(
                _cell(sheet, row, headers, "约束级别"), FORMULA_ENFORCEMENT_ALIASES,
                f"公式“{formula_id}”约束级别", "hard",
            ),
            auto_complete=_bool(_cell(sheet, row, headers, "生成时自动回算"), default=True),
            description=_text(_cell(sheet, row, headers, "说明")),
            expression=expression,
            relation=(
                _enum(
                    _cell(sheet, row, headers, "关系类型"), FORMULA_RELATION_ALIASES,
                    f"公式“{formula_id}”关系类型", "equal",
                )
                if is_v5 else "equal"
            ),
            operating_condition=_text(_cell(sheet, row, headers, "适用工况")) if is_v5 else "",
        )
        if spec.auto_complete and spec.relation != "equal":
            raise ProjectDataError(f"公式“{formula_id}”是非等式关系，不能设置生成时自动回算。")
        if spec.auto_complete and target.participates_generation:
            warnings.append(
                f"派生属性“{target.label}”已设置生成时自动回算，建议在“属性配置”中把“参与生成”改为“否”。"
            )
        formulas.append(spec)
        seen_ids.add(formula_id)
        seen_targets.add(target.key)
    return formulas, warnings


def _read_system_attributes(workbook: Any) -> List[SystemAttributeSpec]:
    if SYSTEM_ATTRIBUTE_SHEET not in workbook.sheetnames:
        raise ProjectDataError(f"缺少“{SYSTEM_ATTRIBUTE_SHEET}”工作表。")
    sheet = workbook[SYSTEM_ATTRIBUTE_SHEET]
    if not _has_headers(sheet, SYSTEM_ATTRIBUTE_HEADERS):
        raise ProjectDataError(f"“{SYSTEM_ATTRIBUTE_SHEET}”表头不完整，请使用V11模板。")
    headers = _header_map(sheet)
    output: List[SystemAttributeSpec] = []
    seen: set[str] = set()
    labels: set[str] = set()
    for row in range(2, sheet.max_row + 1):
        label = _text(_cell(sheet, row, headers, "系统属性"))
        key_raw = _text(_cell(sheet, row, headers, "系统属性ID"))
        if not label and not key_raw:
            continue
        if not label or not key_raw:
            raise ProjectDataError(f"“{SYSTEM_ATTRIBUTE_SHEET}”第{row}行必须同时填写系统属性和系统属性ID。")
        key = _stable_key(key_raw, label, "系统属性")
        if key in seen or _normal(label) in labels:
            raise ProjectDataError(f"系统属性“{label}”或ID“{key}”重复。")
        seen.add(key)
        labels.add(_normal(label))
        output.append(
            SystemAttributeSpec(
                key, label, _text(_cell(sheet, row, headers, "系统单位")) or "-",
                _enum(_cell(sheet, row, headers, "汇总方式"), AGGREGATION_ALIASES, f"第{row}行汇总方式", "single"),
                _text(_cell(sheet, row, headers, "说明")),
            )
        )
    if not output:
        raise ProjectDataError(f"“{SYSTEM_ATTRIBUTE_SHEET}”中至少需要一个系统属性。")
    return output


def _infer_channel_rows(product: ProjectDataset, mappings: List[SystemMappingSpec]) -> List[SystemMappingSpec]:
    covered = {item.product_key for item in mappings}
    next_index = len(mappings) + 1
    for spec in product.attributes:
        if spec.key in covered:
            continue
        channel = "product_performance" if spec.participates_utility else "feasibility_only"
        mappings.append(
            SystemMappingSpec(
                f"AUTO-{next_index:03d}", spec.key, spec.label, spec.unit, channel,
                mapping_type="product_performance" if channel == "product_performance" else "none",
                description="V2根据成品属性配置自动补充的评价通道。",
            )
        )
        next_index += 1
    return mappings


def _read_v2_mappings(workbook: Any, product: ProjectDataset, system_attributes: Sequence[SystemAttributeSpec]) -> List[SystemMappingSpec]:
    if MAPPING_SHEET not in workbook.sheetnames:
        raise ProjectDataError(f"缺少“{MAPPING_SHEET}”工作表。")
    sheet = workbook[MAPPING_SHEET]
    if not _has_headers(sheet, MAPPING_HEADERS):
        raise ProjectDataError(f"“{MAPPING_SHEET}”表头不是V2格式。")
    headers = _header_map(sheet)
    product_by_key = product.attribute_by_key()
    product_by_label = {_normal(item.label): item for item in product.attributes}
    system_by_key = {item.key: item for item in system_attributes}
    output: List[SystemMappingSpec] = []
    seen_ids: set[str] = set()
    role_by_product: Dict[str, set[str]] = {}
    for row in range(2, sheet.max_row + 1):
        mapping_id = _text(_cell(sheet, row, headers, "映射编号"))
        product_key = _text(_cell(sheet, row, headers, "属性ID"))
        product_label = _text(_cell(sheet, row, headers, "成品属性"))
        if not mapping_id and not product_key and not product_label:
            continue
        if not mapping_id:
            mapping_id = f"MAP-{row - 1:03d}"
        if mapping_id in seen_ids:
            raise ProjectDataError(f"映射编号“{mapping_id}”重复。")
        seen_ids.add(mapping_id)
        spec = product_by_key.get(product_key) or product_by_label.get(_normal(product_label))
        if not spec:
            raise ProjectDataError(f"“{MAPPING_SHEET}”第{row}行引用了未知成品属性“{product_label or product_key}”。")
        if product_key and product_key != spec.key:
            raise ProjectDataError(f"第{row}行成品属性名称与属性ID不匹配。")
        channel = _enum(_cell(sheet, row, headers, "评价通道"), CHANNEL_ALIASES, f"第{row}行评价通道", "product_performance")
        roles = role_by_product.setdefault(spec.key, set())
        roles.add(channel)
        if "direct_system" in roles and "product_performance" in roles:
            raise ProjectDataError(f"成品属性“{spec.label}”不能同时完整计入直接系统贡献和成品性能贡献，以免重复计分。")
        system_key = _text(_cell(sheet, row, headers, "系统属性ID"))
        system_attribute = _text(_cell(sheet, row, headers, "系统属性"))
        mapping_type = _enum(_cell(sheet, row, headers, "映射类型"), MAPPING_TYPE_ALIASES, f"第{row}行映射类型", "none")
        coefficient = _optional_float(_cell(sheet, row, headers, "映射系数"), f"第{row}行映射系数")
        constant = _optional_float(_cell(sheet, row, headers, "映射常数"), f"第{row}行映射常数")
        target: Optional[SystemAttributeSpec] = None
        if channel == "direct_system":
            target = system_by_key.get(system_key)
            if not target:
                raise ProjectDataError(f"第{row}行直接映射引用了未知系统属性ID“{system_key}”。")
            if system_attribute and _normal(system_attribute) != _normal(target.label):
                raise ProjectDataError(f"第{row}行系统属性名称与系统属性ID不匹配。")
            if not spec.is_numeric:
                raise ProjectDataError(f"直接系统贡献属性“{spec.label}”必须是数值属性。")
            if mapping_type not in {"linear", "additive"}:
                raise ProjectDataError(f"直接系统映射“{mapping_id}”必须选择线性传导或加总贡献。")
        else:
            system_key = ""
            system_attribute = ""
            mapping_type = "product_performance" if channel == "product_performance" else "none"
        output.append(
            SystemMappingSpec(
                mapping_id, spec.key, spec.label, spec.unit, channel,
                system_key, target.label if target else "", target.unit if target else "",
                mapping_type, 1.0 if coefficient is None else coefficient,
                0.0 if constant is None else constant,
                _text(_cell(sheet, row, headers, "说明")),
            )
        )
    output = _infer_channel_rows(product, output)
    direct_by_system: Dict[str, List[SystemMappingSpec]] = {}
    for item in output:
        if item.channel == "direct_system":
            direct_by_system.setdefault(item.system_key, []).append(item)
    for attr in system_attributes:
        sources = direct_by_system.get(attr.key, [])
        if not sources:
            raise ProjectDataError(f"系统属性“{attr.label}”没有任何成品到系统映射。")
        if attr.aggregation == "single" and len(sources) != 1:
            raise ProjectDataError(f"系统属性“{attr.label}”配置为单值，但存在{len(sources)}条映射；请改为求和或只保留一条。")
    return output


def _validate_requirement(requirement: SystemRequirementSpec, context: str) -> None:
    kind = requirement.requirement_type
    if kind == "target" and requirement.target_value is None:
        raise ProjectDataError(f"{context}的目标值不能为空。")
    if kind == "at_least" and requirement.minimum is None:
        raise ProjectDataError(f"{context}的要求下限不能为空。")
    if kind == "within_range":
        if requirement.minimum is None or requirement.maximum is None:
            raise ProjectDataError(f"{context}必须同时填写要求下限和要求上限。")
        if requirement.minimum >= requirement.maximum:
            raise ProjectDataError(f"{context}的要求下限必须小于要求上限。")
    if kind == "at_most" and requirement.maximum is None:
        return
    for value, label in (
        (requirement.lower_80_distance, "下偏80分距离"),
        (requirement.upper_80_distance, "上偏80分距离"),
        (requirement.lower_compliance_tolerance, "达标允许下偏差"),
        (requirement.upper_compliance_tolerance, "达标允许上偏差"),
    ):
        if value is not None and value < 0:
            raise ProjectDataError(f"{context}的{label}不能为负数。")
        if value == 0 and "80分距离" in label:
            raise ProjectDataError(f"{context}的{label}必须大于0。")
    if requirement.score_cap < 100:
        raise ProjectDataError(f"{context}的单项分数上限不能低于100。")
    if requirement.weight <= 0:
        raise ProjectDataError(f"{context}的组内权重必须大于0。")


def _automatic_distance_description(anchor: Optional[float], unit: str) -> str:
    if anchor is not None and abs(float(anchor)) > 1e-12:
        value = abs(float(anchor)) * 0.10
        return f"自动采用要求值绝对值的10%，即 {value:g} {unit}".strip()
    return "自动根据历史样本的稳健离散尺度计算"


def _requirement_distance_warnings(tasks: Sequence[SystemTask]) -> List[str]:
    """Make every implicit 80-point scale visible without rejecting old workbooks."""
    output: List[str] = []
    for task in tasks:
        for requirement in task.requirements:
            context = f"任务“{task.task_name}”的“{requirement.system_attribute}”"
            kind = requirement.requirement_type
            missing: List[Tuple[str, Optional[float], str]] = []
            if kind == "target":
                missing.extend(
                    [
                        ("下偏80分距离", requirement.target_value, "低于目标时的扣分尺度"),
                        ("上偏80分距离", requirement.target_value, "高于目标时的扣分尺度"),
                    ]
                )
            elif kind == "at_least":
                missing.append(("下偏80分距离", requirement.minimum, "低于下限时的扣分尺度"))
                if requirement.overachievement_bonus:
                    missing.append(("上偏80分距离", requirement.minimum, "高于下限时的加分尺度"))
            elif kind == "at_most" and requirement.maximum is not None:
                missing.append(("上偏80分距离", requirement.maximum, "高于上限时的扣分尺度"))
                if requirement.overachievement_bonus:
                    missing.append(("下偏80分距离", requirement.maximum, "低于上限时的加分尺度"))
            elif kind == "within_range":
                missing.extend(
                    [
                        ("下偏80分距离", requirement.minimum, "低于区间时的扣分尺度"),
                        ("上偏80分距离", requirement.maximum, "高于区间时的扣分尺度"),
                    ]
                )
            for field_name, anchor, meaning in missing:
                configured = (
                    requirement.lower_80_distance
                    if field_name.startswith("下偏")
                    else requirement.upper_80_distance
                )
                if configured is not None:
                    continue
                automatic = _automatic_distance_description(anchor, requirement.system_unit)
                output.append(
                    f"{context}未填写“{field_name}”（{meaning}），{automatic}。"
                    "这是自动默认值，不代表专家确认的工程容差；正式使用前建议在Excel中明确填写。"
                )
    return output


def _read_v2_tasks(workbook: Any, system_attributes: Sequence[SystemAttributeSpec]) -> List[SystemTask]:
    if TASK_SHEET not in workbook.sheetnames:
        raise ProjectDataError(f"缺少“{TASK_SHEET}”工作表。")
    sheet = workbook[TASK_SHEET]
    if not _has_headers(sheet, BASE_TASK_HEADERS):
        raise ProjectDataError(f"“{TASK_SHEET}”表头不是系统任务长表格式。")
    headers = _header_map(sheet)
    system_by_key = {item.key: item for item in system_attributes}
    groups: Dict[str, Dict[str, Any]] = {}
    order: List[str] = []
    for row in range(2, sheet.max_row + 1):
        task_id = _text(_cell(sheet, row, headers, "任务编号"))
        task_name = _text(_cell(sheet, row, headers, "任务名称"))
        system_key = _text(_cell(sheet, row, headers, "系统属性ID"))
        if not task_id and not task_name and not system_key:
            continue
        if not task_id or not task_name:
            raise ProjectDataError(f"“{TASK_SHEET}”第{row}行必须同时填写任务编号和任务名称。")
        group = groups.get(task_id)
        if group is None:
            group = {
                "name": task_name,
                "condition": _text(_cell(sheet, row, headers, "系统型号/工况")),
                "description": _text(_cell(sheet, row, headers, "说明")),
                "requirements": [],
                "keys": set(),
            }
            groups[task_id] = group
            order.append(task_id)
        elif task_name != group["name"]:
            raise ProjectDataError(f"任务编号“{task_id}”对应了多个任务名称。")
        if not system_key:
            continue
        attr = system_by_key.get(system_key)
        if not attr:
            raise ProjectDataError(f"第{row}行引用了未知系统属性ID“{system_key}”。")
        if system_key in group["keys"]:
            raise ProjectDataError(f"任务“{task_name}”中的系统属性“{attr.label}”重复。")
        group["keys"].add(system_key)
        written_label = _text(_cell(sheet, row, headers, "系统属性"))
        if written_label and _normal(written_label) != _normal(attr.label):
            raise ProjectDataError(f"第{row}行系统属性名称与系统属性ID不匹配。")
        score_cap = _optional_float(_cell(sheet, row, headers, "单项分数上限"), f"第{row}行单项分数上限")
        weight = _optional_float(_cell(sheet, row, headers, "组内权重"), f"第{row}行组内权重")
        requirement = SystemRequirementSpec(
            attr.key, attr.label, attr.unit,
            _enum(_cell(sheet, row, headers, "要求类型"), REQUIREMENT_TYPE_ALIASES, f"第{row}行要求类型", "target"),
            _optional_float(_cell(sheet, row, headers, "目标值"), f"第{row}行目标值"),
            _optional_float(_cell(sheet, row, headers, "要求下限"), f"第{row}行要求下限"),
            _optional_float(_cell(sheet, row, headers, "要求上限"), f"第{row}行要求上限"),
            _optional_float(_cell(sheet, row, headers, "下偏80分距离"), f"第{row}行下偏80分距离"),
            _optional_float(_cell(sheet, row, headers, "上偏80分距离"), f"第{row}行上偏80分距离"),
            _bool(_cell(sheet, row, headers, "允许超额加分")),
            100.0 if score_cap is None else score_cap,
            1.0 if weight is None else weight,
            _enum(_cell(sheet, row, headers, "硬性策略"), HARD_POLICY_ALIASES, f"第{row}行硬性策略", "none"),
            _text(_cell(sheet, row, headers, "说明")),
            lower_compliance_tolerance=_optional_float(
                _cell(sheet, row, headers, "达标允许下偏差"), f"第{row}行达标允许下偏差"
            ),
            upper_compliance_tolerance=_optional_float(
                _cell(sheet, row, headers, "达标允许上偏差"), f"第{row}行达标允许上偏差"
            ),
        )
        _validate_requirement(requirement, f"任务“{task_name}”的“{attr.label}”")
        group["requirements"].append(requirement)
    tasks: List[SystemTask] = []
    required_keys = set(system_by_key)
    for task_id in order:
        group = groups[task_id]
        missing = required_keys - set(group["keys"])
        if missing:
            labels = "、".join(system_by_key[key].label for key in sorted(missing))
            raise ProjectDataError(f"任务“{group['name']}”缺少系统属性要求：{labels}。")
        tasks.append(
            SystemTask(
                task_id, group["name"], group["condition"], tuple(group["requirements"]), group["description"]
            )
        )
    if not tasks:
        raise ProjectDataError(f"“{TASK_SHEET}”中至少需要一个系统任务。")
    return tasks


def _legacy_system_key(label: str, index: int) -> str:
    known = {
        "系统压力": "system_pressure",
        "系统流量": "system_flow",
        "泵增重贡献": "system_weight_contribution",
        "系统重量贡献": "system_weight_contribution",
    }
    return known.get(label) or _generated_key(label, "system_attr", index)


def _read_legacy_contract(workbook: Any, product: ProjectDataset) -> Tuple[List[SystemAttributeSpec], List[SystemMappingSpec], List[SystemTask], List[str]]:
    """Read the fixed pressure/flow/weight S1 sheets into the V2 in-memory model."""
    warnings = ["当前工作簿仍为系统级S1格式，已按V2兼容读取；建议使用升级工具生成V2长表。"]
    if MAPPING_SHEET not in workbook.sheetnames or TASK_SHEET not in workbook.sheetnames:
        raise ProjectDataError("Excel缺少系统级V2工作表，请先执行升级工具。")
    mapping_sheet = workbook[MAPPING_SHEET]
    headers = _header_map(mapping_sheet)
    product_by_key = product.attribute_by_key()
    product_by_label = {_normal(item.label): item for item in product.attributes}
    attributes: List[SystemAttributeSpec] = []
    attr_by_label: Dict[str, SystemAttributeSpec] = {}
    mappings: List[SystemMappingSpec] = []
    legacy_requirement_settings: Dict[str, Dict[str, Any]] = {}
    for row in range(2, mapping_sheet.max_row + 1):
        product_key = _text(_cell(mapping_sheet, row, headers, "属性ID"))
        product_label = _text(_cell(mapping_sheet, row, headers, "成品属性"))
        if not product_key and not product_label:
            continue
        spec = product_by_key.get(product_key) or product_by_label.get(_normal(product_label))
        if not spec:
            raise ProjectDataError(f"旧系统映射第{row}行引用了未知成品属性。")
        channel = _enum(_cell(mapping_sheet, row, headers, "评价通道"), CHANNEL_ALIASES, f"第{row}行评价通道", "product_performance")
        system_label = _text(_cell(mapping_sheet, row, headers, "系统属性"))
        system_key = ""
        target: Optional[SystemAttributeSpec] = None
        if channel == "direct_system":
            if not system_label:
                raise ProjectDataError(f"旧系统映射第{row}行缺少系统属性。")
            target = attr_by_label.get(_normal(system_label))
            if target is None:
                system_key = _legacy_system_key(system_label, len(attributes) + 1)
                aggregation = "sum" if _normal(_cell(mapping_sheet, row, headers, "映射类型")) in {"加总贡献", "additive"} else "single"
                target = SystemAttributeSpec(system_key, system_label, _text(_cell(mapping_sheet, row, headers, "系统单位")) or spec.unit, aggregation)
                attributes.append(target)
                attr_by_label[_normal(system_label)] = target
            system_key = target.key
            requirement_type = _enum(_cell(mapping_sheet, row, headers, "要求类型"), REQUIREMENT_TYPE_ALIASES, f"第{row}行要求类型", "target")
            legacy_requirement_settings[system_key] = {
                "type": requirement_type,
                "bonus": _bool(_cell(mapping_sheet, row, headers, "允许超额加分")),
                "cap": _optional_float(_cell(mapping_sheet, row, headers, "单项分数上限"), f"第{row}行单项分数上限") or 100.0,
                "hard": _enum(_cell(mapping_sheet, row, headers, "硬性策略"), HARD_POLICY_ALIASES, f"第{row}行硬性策略", "none"),
                "description": _text(_cell(mapping_sheet, row, headers, "说明")),
            }
        legacy_coefficient = _optional_float(
            _cell(mapping_sheet, row, headers, "映射系数"), f"第{row}行映射系数"
        )
        legacy_constant = _optional_float(
            _cell(mapping_sheet, row, headers, "映射常数"), f"第{row}行映射常数"
        )
        mappings.append(
            SystemMappingSpec(
                f"LEGACY-{row - 1:03d}", spec.key, spec.label, spec.unit, channel,
                system_key, target.label if target else "", target.unit if target else "",
                "additive" if target and target.aggregation == "sum" else "linear" if target else "product_performance" if channel == "product_performance" else "none",
                1.0 if legacy_coefficient is None else legacy_coefficient,
                0.0 if legacy_constant is None else legacy_constant,
                _text(_cell(mapping_sheet, row, headers, "说明")),
            )
        )
    mappings = _infer_channel_rows(product, mappings)

    task_sheet = workbook[TASK_SHEET]
    task_headers = _header_map(task_sheet)
    tasks: List[SystemTask] = []
    for row in range(2, task_sheet.max_row + 1):
        task_id = _text(_cell(task_sheet, row, task_headers, "任务编号"))
        task_name = _text(_cell(task_sheet, row, task_headers, "任务名称"))
        if not task_id and not task_name:
            continue
        requirements: List[SystemRequirementSpec] = []
        for attr in attributes:
            settings = legacy_requirement_settings[attr.key]
            kind = settings["type"]
            target_value = _optional_float(_cell(task_sheet, row, task_headers, "压力需求点"), "压力需求点") if kind == "target" else None
            minimum = _optional_float(_cell(task_sheet, row, task_headers, "流量要求下限"), "流量要求下限") if kind == "at_least" else None
            maximum = _optional_float(_cell(task_sheet, row, task_headers, "重量预算"), "重量预算") if kind == "at_most" else None
            requirement = SystemRequirementSpec(
                attr.key, attr.label, attr.unit, kind, target_value, minimum, maximum,
                _optional_float(_cell(task_sheet, row, task_headers, "压力允许下偏"), "压力允许下偏") if kind == "target" else None,
                _optional_float(_cell(task_sheet, row, task_headers, "压力允许上偏"), "压力允许上偏") if kind == "target" else None,
                settings["bonus"], settings["cap"], 1.0, settings["hard"], settings["description"],
            )
            _validate_requirement(requirement, f"任务“{task_name}”的“{attr.label}”")
            requirements.append(requirement)
        tasks.append(
            SystemTask(
                task_id, task_name, _text(_cell(task_sheet, row, task_headers, "系统型号/工况")),
                tuple(requirements), _text(_cell(task_sheet, row, task_headers, "说明")),
            )
        )
    if not tasks:
        raise ProjectDataError("旧系统任务表中没有可用任务。")
    return attributes, mappings, tasks, warnings


def _system_fingerprint(
    product: ProjectDataset,
    attributes: Sequence[SystemAttributeSpec],
    mappings: Sequence[SystemMappingSpec],
    tasks: Sequence[SystemTask],
    formulas: Sequence[FormulaConstraintSpec],
) -> str:
    payload = {
        "version": CONTRACT_VERSION,
        "product_learning_fingerprint": product.learning_fingerprint,
        "system_attributes": [asdict(item) for item in attributes],
        "mappings": [asdict(item) for item in mappings],
        "tasks": [asdict(item) for item in tasks],
        "formulas": [asdict(item) for item in formulas],
    }
    serialized = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(serialized.encode("utf-8")).hexdigest()[:20]


def _system_learning_fingerprint(
    product: ProjectDataset,
    attributes: Sequence[SystemAttributeSpec],
    mappings: Sequence[SystemMappingSpec],
    formulas: Sequence[FormulaConstraintSpec],
) -> str:
    """Identify reusable system knowledge without binding it to task targets."""
    payload = {
        "version": "system_learning_contract_v11_requirement_anchored_weight_separated_frontier",
        "product_learning_fingerprint": product.learning_fingerprint,
        "system_attributes": [asdict(item) for item in attributes],
        "mappings": [asdict(item) for item in mappings],
        "formulas": [asdict(item) for item in formulas],
    }
    serialized = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(serialized.encode("utf-8")).hexdigest()[:20]


def _is_v2_workbook(workbook: Any) -> bool:
    return (
        SYSTEM_ATTRIBUTE_SHEET in workbook.sheetnames
        and MAPPING_SHEET in workbook.sheetnames
        and TASK_SHEET in workbook.sheetnames
        and _has_headers(workbook[MAPPING_SHEET], MAPPING_HEADERS)
        and _has_headers(workbook[TASK_SHEET], BASE_TASK_HEADERS)
    )


def load_system_workbook(path: Union[Path, str]) -> SystemProjectDataset:
    workbook_path = Path(path).expanduser().resolve()
    product = load_project_workbook(workbook_path)
    workbook = load_workbook(BytesIO(workbook_path.read_bytes()), data_only=True, read_only=True)
    try:
        if _is_v2_workbook(workbook):
            attributes = _read_system_attributes(workbook)
            mappings = _read_v2_mappings(workbook, product, attributes)
            tasks = _read_v2_tasks(workbook, attributes)
            formulas, formula_warnings = _read_formulas(workbook, product)
            warnings: List[str] = list(formula_warnings)
            if FORMULA_SHEET not in workbook.sheetnames:
                source_version = "system_excel_v3_compat"
            elif _has_headers(workbook[FORMULA_SHEET], FORMULA_HEADERS):
                source_version = CONTRACT_VERSION
            else:
                source_version = "system_excel_v4_compat"
        else:
            attributes, mappings, tasks, warnings = _read_legacy_contract(workbook, product)
            formulas = []
            warnings.append("旧版系统级Excel已兼容读取，但尚未配置派生属性公式。")
            source_version = "system_excel_v1_compat"
    finally:
        workbook.close()
    if not [item for item in mappings if item.channel == "direct_system"]:
        raise ProjectDataError("系统级项目至少需要一条直接系统贡献映射。")
    product_group = [item for item in mappings if item.channel == "product_performance"]
    if not product_group:
        warnings.append("当前没有成品性能贡献属性，综合分将完全由系统要求贡献组成。")
    warnings.extend(_requirement_distance_warnings(tasks))
    return SystemProjectDataset(
        product=product,
        system_attributes=attributes,
        mappings=mappings,
        tasks=tasks,
        system_fingerprint=_system_fingerprint(product, attributes, mappings, tasks, formulas),
        system_learning_fingerprint=_system_learning_fingerprint(product, attributes, mappings, formulas),
        formulas=formulas,
        contract_version=CONTRACT_VERSION,
        source_contract_version=source_version,
        warnings=warnings,
    )


def upgrade_product_workbook(source: Union[Path, str], destination: Union[Path, str]) -> Path:
    source_path = Path(source).expanduser().resolve()
    destination_path = Path(destination).expanduser().resolve()
    if source_path == destination_path:
        raise ProjectDataError("升级输出必须是新文件，不能覆盖原Excel。")
    if not source_path.is_file():
        raise ProjectDataError(f"源Excel不存在：{source_path}")
    destination_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(str(source_path), str(destination_path))
    product = load_project_workbook(source_path)
    source_workbook = load_workbook(BytesIO(source_path.read_bytes()), data_only=True, read_only=True)
    formulas: List[FormulaConstraintSpec] = []
    try:
        is_current_base = _is_v2_workbook(source_workbook)
        if is_current_base:
            formulas, _ = _read_formulas(source_workbook, product)
        if not is_current_base and MAPPING_SHEET in source_workbook.sheetnames and TASK_SHEET in source_workbook.sheetnames:
            attributes, mappings, tasks, _ = _read_legacy_contract(source_workbook, product)
        elif not is_current_base:
            attributes, mappings, tasks = _default_contract(product)
    finally:
        source_workbook.close()
    workbook = load_workbook(destination_path)
    try:
        if is_current_base:
            _write_formula_sheet(workbook, formulas)
        else:
            _write_v2_contract(workbook, attributes, mappings, tasks)
        workbook.save(destination_path)
    finally:
        workbook.close()
    return destination_path


__all__ = [
    "CONTRACT_VERSION",
    "SYSTEM_ATTRIBUTE_SHEET",
    "MAPPING_SHEET",
    "TASK_SHEET",
    "FORMULA_SHEET",
    "SystemAttributeSpec",
    "SystemMappingSpec",
    "SystemRequirementSpec",
    "SystemTask",
    "SystemProjectDataset",
    "FormulaConstraintSpec",
    "detect_direct_attributes",
    "load_system_workbook",
    "upgrade_product_workbook",
]
